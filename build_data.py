"""Genera data.js da 'Matrici SNA.xlsx'. Da rilanciare solo se cambiano i dati."""

import json
import networkx as nx
from openpyxl import load_workbook

XLSX = "Matrici SNA.xlsx"
GRAPHS = ["ALL", "GG-EU", "GG-CH"]
N = 48
SPREAD = 900  # semi-lato in px del riquadro in cui distribuire i nodi

wb = load_workbook(XLSX, read_only=True, data_only=True)

# Foglio Topics: indice -> nome del concetto.
labels = {
    int(i): str(name)
    for i, name in wb["Topics"].iter_rows(max_row=N, max_col=2, values_only=True)
}


def read_matrix(sheet):
    """Matrice NxN di interi, saltando la riga e la colonna degli indici."""
    rows = sheet.iter_rows(min_row=2, max_row=N + 1, min_col=2, max_col=N + 1, values_only=True)
    return [[int(v or 0) for v in row] for row in rows]


# Le matrici sono simmetriche: la diagonale e' il peso del nodo, il resto quello
# dell'arco. Per gli archi basta quindi il triangolo superiore.
graphs = {}
for name in GRAPHS:
    m = read_matrix(wb[name])
    graphs[name] = {
        "weights": {str(i + 1): m[i][i] for i in range(N)},
        "edges": [[i + 1, j + 1, m[i][j]] for i in range(N) for j in range(i + 1, N) if m[i][j]],
    }

# Posizioni calcolate una volta sola su ALL e condivise dai tre grafi, cosi' nello
# switch i nodi restano fermi. Il peso conta: ALL e' molto denso e senza pesare le
# co-occorrenze il layout collassa in una palla informe. Seme fisso => riproducibile.
g = nx.Graph()
g.add_nodes_from(range(1, N + 1))
g.add_weighted_edges_from(graphs["ALL"]["edges"])
pos = nx.spring_layout(g, weight="weight", seed=42, iterations=200)

span = max(max(abs(c) for c in p) for p in pos.values())
nodes = [
    {
        "id": str(i),
        "label": labels[i],
        "x": round(pos[i][0] / span * SPREAD, 1),
        "y": round(pos[i][1] / span * SPREAD, 1),
    }
    for i in range(1, N + 1)
]

# Un nodo per riga e un grafo per riga: e' un file generato, ma i diff restano leggibili.
nodes_js = ",\n  ".join(json.dumps(n, ensure_ascii=False) for n in nodes)
graphs_js = ",\n  ".join(f"{json.dumps(k)}: {json.dumps(v)}" for k, v in graphs.items())

with open("data.js", "w", encoding="utf-8") as f:
    f.write("// Generato da build_data.py - non modificare a mano.\n")
    f.write(f"const DATA = {{\n nodes: [\n  {nodes_js}\n ],\n graphs: {{\n  {graphs_js}\n }}\n}};\n")

print(f"data.js: {len(nodes)} nodi, " + ", ".join(f"{k} {len(v['edges'])} archi" for k, v in graphs.items()))
